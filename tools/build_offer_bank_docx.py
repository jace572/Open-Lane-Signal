#!/usr/bin/env python3
"""
Convert an offer bank markdown file into a polished .docx deliverable.

Usage:
    python3 build_offer_bank_docx.py INPUT.md [OUTPUT.docx]

If OUTPUT.docx is omitted, it's saved next to the input with .docx extension.

Recognized markdown patterns:
    # Title              → H1
    ## Section            → H2
    ### Item              → H3
    > "quote text"        → indented italic blockquote with grey left border
    > — Attribution       → small italic attribution line below blockquote
    **Label:** body       → bold-label paragraph (Why it works:, Best for:, etc.)
    - Item                → bullet list
    1. Item               → numbered list
    | a | b |             → table (with --- separator row)
    ---                   → horizontal rule
    plain text            → paragraph

Run from the SOPs/tools directory or pass absolute paths.
Requires python-docx: pip3 install python-docx
"""

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


# ---------------- Style constants ----------------
H1_COLOR = RGBColor(0x1F, 0x2A, 0x44)
H2_COLOR = RGBColor(0x1F, 0x2A, 0x44)
H3_COLOR = RGBColor(0x2E, 0x4A, 0x6B)
BLOCKQUOTE_TEXT_COLOR = RGBColor(0x44, 0x55, 0x66)
BLOCKQUOTE_BORDER_HEX = "BFC8D1"
ATTRIBUTION_COLOR = RGBColor(0x66, 0x77, 0x88)
TABLE_HEADER_BG = "2E4A6B"
TABLE_HEADER_TEXT = RGBColor(0xFF, 0xFF, 0xFF)
RULE_COLOR = "C0C8D0"


# ---------------- Document setup ----------------
def make_doc():
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    style_normal = doc.styles["Normal"]
    style_normal.font.name = "Calibri"
    style_normal.font.size = Pt(11)

    def set_heading(name, size_pt, color, space_before, space_after):
        s = doc.styles[name]
        s.font.name = "Calibri"
        s.font.size = Pt(size_pt)
        s.font.bold = True
        s.font.color.rgb = color
        s.paragraph_format.space_before = Pt(space_before)
        s.paragraph_format.space_after = Pt(space_after)
        s.paragraph_format.keep_with_next = True

    set_heading("Heading 1", 22, H1_COLOR, 0, 12)
    set_heading("Heading 2", 16, H2_COLOR, 18, 6)
    set_heading("Heading 3", 13, H3_COLOR, 12, 4)
    return doc


# ---------------- Paragraph builders ----------------
def add_blockquote(doc, quote_text, attribution=None):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.4)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)

    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    left_border = OxmlElement("w:left")
    left_border.set(qn("w:val"), "single")
    left_border.set(qn("w:sz"), "18")
    left_border.set(qn("w:space"), "8")
    left_border.set(qn("w:color"), BLOCKQUOTE_BORDER_HEX)
    pBdr.append(left_border)
    pPr.append(pBdr)

    run = p.add_run(f"“{quote_text}”")
    run.italic = True
    run.font.color.rgb = BLOCKQUOTE_TEXT_COLOR

    if attribution:
        p2 = doc.add_paragraph()
        p2.paragraph_format.left_indent = Inches(0.4)
        p2.paragraph_format.space_before = Pt(0)
        p2.paragraph_format.space_after = Pt(8)
        run2 = p2.add_run(f"— {attribution}")
        run2.font.size = Pt(9)
        run2.font.color.rgb = ATTRIBUTION_COLOR
        run2.italic = True


def add_horizontal_rule(doc):
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), RULE_COLOR)
    pBdr.append(bottom)
    pPr.append(pBdr)
    p.paragraph_format.space_after = Pt(8)


def shade_cell(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Light Grid Accent 1"
    table.autofit = True

    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        para = cell.paragraphs[0]
        run = para.add_run(h)
        run.bold = True
        run.font.color.rgb = TABLE_HEADER_TEXT
        run.font.size = Pt(11)
        shade_cell(cell, TABLE_HEADER_BG)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    for r_idx, row_data in enumerate(rows):
        for c_idx, val in enumerate(row_data):
            cell = table.rows[r_idx + 1].cells[c_idx]
            cell.text = ""
            para = cell.paragraphs[0]
            apply_inline(para, val)
            for run in para.runs:
                run.font.size = Pt(10.5)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP


# ---------------- Inline formatting (bold, italic, code) ----------------
INLINE_TOKEN_RE = re.compile(
    r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)"
)

def apply_inline(paragraph, text):
    """Render markdown inline bold/italic/code into a paragraph."""
    parts = INLINE_TOKEN_RE.split(text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("*") and part.endswith("*"):
            run = paragraph.add_run(part[1:-1])
            run.italic = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
        else:
            paragraph.add_run(part)


def add_paragraph_md(doc, text, italic=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    apply_inline(p, text)
    if italic:
        for run in p.runs:
            run.italic = True


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(2)
    apply_inline(p, text)


def add_numbered(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(2)
    apply_inline(p, text)


# ---------------- Markdown parser ----------------
TABLE_SEP_RE = re.compile(r"^\s*\|?\s*[-:|\s]+\s*\|?\s*$")

def parse_table_row(line):
    """Split a markdown table row by | respecting leading/trailing pipes."""
    raw = line.strip()
    if raw.startswith("|"):
        raw = raw[1:]
    if raw.endswith("|"):
        raw = raw[:-1]
    return [cell.strip() for cell in raw.split("|")]


def is_table_separator(line):
    return bool(TABLE_SEP_RE.match(line)) and "|" in line and "-" in line


def convert_md_to_docx(input_path, output_path):
    md_text = Path(input_path).read_text(encoding="utf-8")
    lines = md_text.splitlines()
    doc = make_doc()

    i = 0
    n = len(lines)

    pending_blockquote = None  # holds quote text waiting for attribution

    def flush_blockquote():
        nonlocal pending_blockquote
        if pending_blockquote is not None:
            add_blockquote(doc, pending_blockquote, None)
            pending_blockquote = None

    while i < n:
        line = lines[i]
        stripped = line.strip()

        # Blank line
        if not stripped:
            flush_blockquote()
            i += 1
            continue

        # Horizontal rule
        if stripped == "---" or stripped == "***":
            flush_blockquote()
            add_horizontal_rule(doc)
            i += 1
            continue

        # Headings
        h_match = re.match(r"^(#{1,6})\s+(.+)$", line)
        if h_match:
            flush_blockquote()
            level = len(h_match.group(1))
            text = h_match.group(2).strip()
            doc.add_heading(text, level=min(level, 3))
            i += 1
            continue

        # Blockquote
        if stripped.startswith(">"):
            content = stripped[1:].strip()
            # Attribution line: starts with em-dash or "--"
            if content.startswith("—") or content.startswith("--"):
                attribution = content.lstrip("—- ").strip()
                if pending_blockquote is not None:
                    # Strip surrounding quotes from quote text
                    qt = pending_blockquote.strip()
                    if (qt.startswith('"') and qt.endswith('"')) or \
                       (qt.startswith("“") and qt.endswith("”")):
                        qt = qt[1:-1]
                    add_blockquote(doc, qt, attribution)
                    pending_blockquote = None
                else:
                    add_paragraph_md(doc, content)
            else:
                # Quote line — accumulate. Strip surrounding quotes if present.
                qt = content
                if (qt.startswith('"') and qt.endswith('"')) or \
                   (qt.startswith("“") and qt.endswith("”")):
                    qt = qt[1:-1]
                if pending_blockquote is None:
                    pending_blockquote = qt
                else:
                    pending_blockquote += " " + qt
            i += 1
            continue
        else:
            flush_blockquote()

        # Tables
        if "|" in stripped and i + 1 < n and is_table_separator(lines[i + 1]):
            headers = parse_table_row(line)
            i += 2  # skip separator
            rows = []
            while i < n and "|" in lines[i].strip():
                if not lines[i].strip():
                    break
                rows.append(parse_table_row(lines[i]))
                i += 1
            add_table(doc, headers, rows)
            continue

        # Bullets
        if re.match(r"^\s*[-*+]\s+", line):
            text = re.sub(r"^\s*[-*+]\s+", "", line)
            add_bullet(doc, text)
            i += 1
            continue

        # Numbered list
        if re.match(r"^\s*\d+\.\s+", line):
            text = re.sub(r"^\s*\d+\.\s+", "", line)
            add_numbered(doc, text)
            i += 1
            continue

        # Regular paragraph
        add_paragraph_md(doc, line)
        i += 1

    flush_blockquote()

    doc.save(output_path)
    print(f"Saved: {output_path}")


# ---------------- CLI ----------------
def main():
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(1)

    input_path = Path(sys.argv[1]).expanduser().resolve()
    if not input_path.exists():
        print(f"Error: {input_path} not found")
        sys.exit(1)

    if len(sys.argv) >= 3:
        output_path = Path(sys.argv[2]).expanduser().resolve()
    else:
        output_path = input_path.with_suffix(".docx")

    convert_md_to_docx(input_path, output_path)


if __name__ == "__main__":
    main()
